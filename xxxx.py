"""江苏省大学生安全知识竞赛的协议客户端。

依赖：requests、beautifulsoup4、openpyxl、ddddocr。
默认只读取项目与活动；加入 --execute 才会创建答卷、记录答案并交卷。
密码通过 MARS_PASSWORD 环境变量或终端隐藏输入提供，不写入文件。
"""

from __future__ import annotations

import argparse
import getpass
import hashlib
import json
import os
import random
import re
import sys
import time
import traceback
import unicodedata
import urllib.request
from pathlib import Path
from typing import Any, Callable

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook


BASE_URL = "https://mars.mycourse.cn/marsapi"
SITE_URL = "https://mars.mycourse.cn/deimos/"
ACTIVITY_TYPES = {"sequential": 1, "random": 2, "mock": 3, "formal": 4}
PRACTICE_API = "/api/study/activity/practice/v1"
EXAM_API = "/api/study/activity/exam/v1"


def bundled_bank_path() -> Path:
    """开发环境读取同目录题库；PyInstaller 单文件运行时读取内嵌题库。"""
    return Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent)) / "题库.xlsx"


class PlatformError(RuntimeError):
    def __init__(self, status: Any, message: str):
        self.status = status
        super().__init__(f"平台返回 {status}: {message}")


def clean_text(value: Any) -> str:
    """使用 bs4 解析题目和选项中的 HTML，再统一字符用于精确匹配。"""
    text = BeautifulSoup(str(value or ""), "html.parser").get_text(" ", strip=True)
    text = unicodedata.normalize("NFKC", text)
    return re.sub(r"\s+", "", text).strip()


def decode_result(payload: dict[str, Any]) -> Any:
    if int(payload.get("status", -1)) != 200:
        raise PlatformError(payload.get("status"), str(payload.get("message", "未知错误")))
    result = (payload.get("data") or {}).get("result")
    if isinstance(result, str):
        try:
            return json.loads(result)
        except ValueError:
            return result
    return result


class AnswerResolver:
    def __init__(self, bank_path: Path):
        book = load_workbook(bank_path, read_only=True, data_only=True)
        try:
            sheet = book.active
            headers = list(next(sheet.values))
            if headers[:6] != ["序号", "题型", "类别", "题目", "选项", "答案"]:
                raise ValueError("题库列格式与预期不符")
            self.bank: dict[str, tuple[Any, ...]] = {}
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if row[3] and row[5]:
                    key = clean_text(row[3])
                    if key in self.bank:
                        raise ValueError(f"题库有重复题目：{row[3]}")
                    self.bank[key] = row
        finally:
            book.close()
        self.server_answers: dict[str, set[str]] = {}

    def add_server_questions(self, questions: list[dict[str, Any]]) -> None:
        for question in questions:
            correct = {
                clean_text(option.get("content"))
                for option in question.get("optionList", [])
                if option.get("isCorrect") == 1
            }
            if correct:
                self.server_answers[clean_text(question.get("title"))] = correct

    def _bank_answer_text(self, question: dict[str, Any]) -> set[str]:
        title = clean_text(question.get("title"))
        if title not in self.bank:
            raise ValueError(f"题库未匹配题目：{question.get('title')}")
        row = self.bank[title]
        answer = str(row[5]).strip().upper()
        if int(question.get("type", 0)) == 4:
            if answer == "Y":
                return {"正确", "对", "是", "√"}
            if answer == "N":
                return {"错误", "错", "否", "×"}
            raise ValueError(f"判断题答案格式异常：{answer}")
        bank_options = str(row[4]).split("|")
        labels = [piece.strip() for piece in answer.split("|")]
        if any(len(label) != 1 or label < "A" or label > "Z" for label in labels):
            raise ValueError(f"题库答案格式异常：{answer}")
        indexes = [ord(label) - ord("A") for label in labels]
        if any(index >= len(bank_options) for index in indexes):
            raise ValueError(f"题库选项越界：{answer}")
        return {clean_text(bank_options[index]) for index in indexes}

    def correct_option_ids(self, question: dict[str, Any]) -> set[str]:
        options = question.get("optionList") or []
        if not options:
            raise ValueError(f"题目无选项：{question.get('title')}")
        title = clean_text(question.get("title"))
        desired = self.server_answers.get(title)
        if desired is None:
            desired = self._bank_answer_text(question)
        matches = {
            str(option["id"])
            for option in options
            if clean_text(option.get("content")) in desired
        }
        required = len(desired) if int(question.get("type", 0)) != 4 else 1
        if len(matches) != required:
            raise ValueError(
                f"题目与选项无法可靠对应：{question.get('title')}；"
                f"匹配 {len(matches)}/{required}"
            )
        if int(question.get("type", 0)) in (1, 4) and len(matches) != 1:
            raise ValueError(f"单选或判断题匹配了多个答案：{question.get('title')}")
        return matches


class API:
    """沿用 test.py 的 API 类与 apiMain 调用方式。"""

    def __init__(
        self,
        t: str,
        s: str,
        u: str,
        p: str,
        proxy: str | None,
        t_name: str,
        k_sess: requests.Session,
        bank_path: Path,
        captcha_attempts: int = 3,
    ):
        self.t, self.s, self.u, self.p = t, s, u, p
        self.type_p = t_name
        self.session = k_sess
        self.proxy = proxy
        self.proxies = {"http": proxy, "https": proxy} if proxy else None
        if proxy:
            self.session.trust_env = False
        self.bank_path = bank_path.resolve()
        self.bank = AnswerResolver(self.bank_path)
        self.captcha_attempts = captcha_attempts
        self.coursesList: list[dict[str, Any]] = []
        self.courseInfo: list[dict[str, Any]] = []
        self.c_info: dict[str, Any] | None = None
        self.c = ""
        self.c_id = ""
        self.jds: dict[str, str] = {"jd": "", "jds": ""}
        self.project: dict[str, Any] = {}
        self.activities: dict[int, dict[str, Any]] = {}
        self.login_info: dict[str, Any] = {}
        self.h: dict[str, str] = {}

    def setInfo(self, c: str, c_id: str) -> None:
        self.c, self.c_id = c, c_id

    def setHeader(self) -> None:
        self.h = {
            "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
            "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36",
            "Content-Type": "application/x-www-form-urlencoded;charset=utf8",
        }
        self.session.headers.update(self.h)

    def go_post(self, url: str, headers: dict[str, str] | None, data: dict[str, Any], timeout: int = 25) -> requests.Response:
        return self.session.post(
            url, data=data, headers=headers, timeout=timeout, proxies=self.proxies
        )

    def go_get(self, url: str, headers: dict[str, str] | None, params: dict[str, Any] | None = None, timeout: int = 25) -> requests.Response:
        return self.session.get(
            url, headers=headers, params=params, timeout=timeout, proxies=self.proxies
        )

    def _post(self, path: str, data: dict[str, Any] | None = None, *, relogin: bool = True) -> Any:
        response = self.go_post(BASE_URL + path, None, data or {})
        response.raise_for_status()
        try:
            return decode_result(response.json())
        except PlatformError as error:
            if error.status == 401 and relogin and path != "/api/login/v1/login":
                print("登录态已过期，重新识别验证码并登录")
                if self.login():
                    return self._post(path, data, relogin=False)
            raise

    def update(self, tt: str, state: str, progress: str, detailed: str, remarks: str, z: int = 0) -> None:
        print(f"[update] {z} {tt} {state} {progress} {detailed} {remarks}")

    def my_except(self, except_info: str) -> None:
        print("[my_except]", except_info)

    def _captcha_code(self, timestamp: int, ocr: Any) -> str:
        response = self.go_get(
            BASE_URL + "/api/login/v1/getCaptcha", None, {"timestamp": timestamp}
        )
        response.raise_for_status()
        code = ocr.classification(response.content).strip()
        if not re.fullmatch(r"[A-Za-z0-9]{4,6}", code):
            raise ValueError("ddddocr 未识别出有效验证码")
        return code

    def login(self) -> bool:
        try:
            import ddddocr

            home = self.go_get(SITE_URL, None)
            home.raise_for_status()
            html = BeautifulSoup(home.content, "html.parser")
            if not html.select_one("#app"):
                raise ValueError("目标站首页结构已变化")
            tenants = self._post("/api/login/v1/listTenant")
            tenant = next((item for item in tenants if item.get("name") == self.s), None)
            if tenant is None:
                raise ValueError(f"找不到学校：{self.s}")
            ocr = ddddocr.DdddOcr(show_ad=False)
            for attempt in range(1, self.captcha_attempts + 1):
                timestamp = int(time.time() * 1000) + random.randrange(100000)
                try:
                    code = self._captcha_code(timestamp, ocr)
                except ValueError:
                    print(f"第 {attempt} 次验证码未识别，换图重试")
                    continue
                data = {
                    "tenantCode": tenant["code"],
                    "username": self.u,
                    "encryptedPassword": hashlib.md5(self.p.encode("utf-8")).hexdigest().upper(),
                    "captchaCode": code,
                    "timestamp": timestamp,
                }
                try:
                    self.login_info = self._post("/api/login/v1/login", data, relogin=False)
                except PlatformError as error:
                    if "验证码" in str(error):
                        print(f"第 {attempt} 次验证码校验未通过，换图重试")
                        continue
                    raise
                self.session.headers["Authorization"] = self.login_info["token"]
                if str(self.login_info.get("userName")) != self.u:
                    raise ValueError("登录返回的账号与请求账号不一致")
                return True
            raise ValueError("ddddocr 未能通过验证码校验，已停止登录")
        except Exception:
            self.my_except(traceback.format_exc())
            return False

    def get_coursesList(self) -> bool:
        try:
            self.project = self._post("/api/study/project/v1/index")
            self.activities = {
                int(item["activityType"]): item
                for item in self.project.get("studyActivityList", [])
            }
            self.coursesList = [{
                "courseName": self.project["projectName"],
                "courseId": self.project["projectId"],
                "courseType": "竞赛项目",
                "progress": "",
                "school": self.s,
                "imgUrl": self.project.get("projectImageUrl", ""),
            }]
            return True
        except Exception:
            self.my_except(traceback.format_exc())
            return False

    def getCourses(self) -> bool:
        if self.c_id and self.c_id != "无":
            self.courseInfo = [
                item for item in self.coursesList if self.c_id in str(item["courseId"])
            ]
        else:
            self.courseInfo = [
                item for item in self.coursesList if self.c.strip() in item["courseName"]
            ]
        return bool(self.courseInfo)

    def getDetail(self, detail: dict[str, Any]) -> None:
        self.c_id = str(detail.get("courseId", "无"))
        self.c_info = detail
        self.get_jd()

    def get_jd(self) -> dict[str, str]:
        if self.activities:
            attempted = sum(
                1 for item in self.activities.values()
                if int((item.get("examData") or {}).get("doneTimes", 0)) > 0
            )
            self.jds = {
                "jd": f"{attempted}/{len(self.activities)} 项有答题记录",
                "jds": ", ".join(
                    f"{item['name']} {item.get('examData', {}).get('doneTimes', 0)} 次"
                    for item in self.activities.values()
                ),
            }
        return self.jds

    def handle(self) -> None:
        self.update(self.c, "已检查", self.jds["jd"], "此竞赛项目未提供章节或课件入口", "", 2)

    def list_activities(self) -> None:
        for item in sorted(self.activities.values(), key=lambda x: x["activityType"]):
            data = item.get("examData") or {}
            base = self._api_base(int(item["activityType"]))
            latest = None
            if int(data.get("doneTimes") or 0) > 0:
                try:
                    history = self._post(
                        base + "/listHistory",
                        {"userActivityId": item["userActivityId"]},
                    ) or []
                    latest = history[0] if history else None
                except PlatformError as error:
                    print(f"{item['name']} 历史记录暂不可读：{error}")
            print(
                f"{item['activityType']} {item['name']} "
                f"已答 {data.get('doneTimes', 0)} 次 / 限制 {data.get('totalTimes', -1)} 次 "
                f"最高分 {data.get('maxScore', '未知')}"
                + (
                    f"；最近成绩 {latest.get('score', latest.get('rightNum', '未知'))}，"
                    f"用时 {latest.get('useTime', '未知')} 秒"
                    if latest else ""
                )
            )

    def _api_base(self, activity_type: int) -> str:
        return PRACTICE_API if activity_type in (1, 2) else EXAM_API

    def _answer_history(self) -> None:
        for activity_type in (1, 2):
            activity = self.activities.get(activity_type)
            if not activity:
                continue
            try:
                history = self._post(
                    PRACTICE_API + "/listHistory",
                    {"userActivityId": activity["userActivityId"]},
                ) or []
                if not history:
                    continue
                record = self._post(
                    PRACTICE_API + "/listErrorQuestion.do",
                    {"userExamId": history[0]["id"]},
                ) or {}
                self.bank.add_server_questions(record.get("questionList") or [])
            except PlatformError as error:
                print(f"历史答案接口不可用：{error}；改用本地题库")

    def _choose_wrong(self, question: dict[str, Any], correct_ids: set[str]) -> str | None:
        wrong = [
            str(option["id"])
            for option in question["optionList"]
            if str(option["id"]) not in correct_ids
        ]
        return wrong[0] if wrong else None

    def _score_plan(
        self,
        questions: list[dict[str, Any]],
        answers: dict[str, set[str]],
        paper_score: int,
        target_score: int | None,
        seed: int,
    ) -> set[str]:
        if target_score is None:
            return set()
        if paper_score != len(questions):
            raise ValueError("未验证逐题分值，无法保证目标分数；仅支持总分等于题数的试卷")
        if not 0 <= target_score <= paper_score:
            raise ValueError(f"目标分数须在 0 到 {paper_score} 之间")
        wrong_count = paper_score - target_score
        candidates = [
            question["id"] for question in questions
            if self._choose_wrong(question, answers[question["id"]]) is not None
        ]
        if wrong_count > len(candidates):
            raise ValueError(
                f"只能为 {len(candidates)} 道题构造确定的零分答案；"
                f"目标 {target_score} 分无法可靠实现"
            )
        return set(random.Random(seed).sample(candidates, wrong_count))

    def _record_answer(
        self,
        activity_type: int,
        activity_id: str,
        question: dict[str, Any],
        chosen_ids: set[str],
        current_index: int,
        seconds: int,
    ) -> None:
        option_order = [str(item["id"]) for item in question["optionList"]]
        ordered = [option_id for option_id in option_order if option_id in chosen_ids]
        if int(question["type"]) in (1, 4) and len(ordered) != 1:
            raise ValueError("单选或判断题必须选一个答案")
        params: dict[str, Any] = {
            "userActivityId": activity_id,
            "questionId": question["id"],
            "useTime": max(1, seconds),
            "answerIds": ",".join(ordered),
        }
        if activity_type in (1, 2):
            params["currentIndex"] = current_index
        for attempt in range(4):
            try:
                self._post(self._api_base(activity_type) + "/recordQuestion", params)
                return
            except (requests.exceptions.Timeout, requests.exceptions.ConnectionError):
                if attempt == 3:
                    raise
                print(f"题目 {current_index} 记录超时，重试 {attempt + 1}/3")
                time.sleep(1 + attempt * 2)

    def run_activity(
        self,
        activity_type: int,
        *,
        execute: bool,
        target_score: int | None,
        duration_seconds: int | None,
        seed: int,
    ) -> dict[str, Any] | None:
        activity = self.activities.get(activity_type)
        if activity is None:
            raise ValueError(f"项目没有活动类型 {activity_type}")
        activity_id = str(activity["userActivityId"])
        base = self._api_base(activity_type)
        preparation = self._post(base + "/preparePaper", {"userActivityId": activity_id})
        print(
            f"{activity['name']}：{preparation.get('questionNum')} 题，"
            f"总分 {preparation.get('paperScore')}，限时 {preparation.get('answerTime')} 分钟，"
            f"已用 {preparation.get('doneTimes', 0)} 次"
        )
        if not execute:
            return preparation
        total_times = int(preparation.get("totalTimes") or -1)
        done_times = int(preparation.get("doneTimes") or 0)
        if total_times > 0 and done_times >= total_times:
            raise ValueError("答题次数已经用完")
        if activity_type in (1, 2) and target_score is not None:
            raise ValueError("练习无分数目标；--target-score 仅适用于模拟或正式考试")
        if duration_seconds is not None:
            limit = int(preparation.get("answerTime") or 0) * 60
            if duration_seconds < 1 or (limit and duration_seconds >= limit):
                raise ValueError(f"目标时长必须介于 1 和 {limit - 1} 秒之间")
        if activity_type in (3, 4):
            score = int(preparation.get("paperScore") or 0)
            count = int(preparation.get("questionNum") or 0)
            if target_score is not None and (score != count or not 0 <= target_score <= score):
                raise ValueError("目标分数不在可验证范围内")
        self._answer_history()

        started_at = time.monotonic()
        paper = self._post(base + "/startPaper", {"userActivityId": activity_id})
        current_index = int(paper.get("currentIndex") or 1)
        questions = paper.get("questionList") or []
        if not questions:
            raise ValueError("取卷结果没有题目；已停止，未交卷")
        if activity_type in (3, 4) and len(questions) != int(preparation["questionNum"]):
            raise ValueError("考试取卷未返回全部题目；已停止，未交卷")

        processed = 0
        wrong_ids: set[str] = set()
        if activity_type in (3, 4):
            answers = {question["id"]: self.bank.correct_option_ids(question) for question in questions}
            wrong_ids = self._score_plan(
                questions, answers, int(preparation["paperScore"]), target_score, seed
            )
            print(
                f"试卷校验：{len(questions)} 题均已匹配答案，"
                f"计划答对 {len(questions) - len(wrong_ids)} 题"
            )
        else:
            answers = {}
        last_record = time.monotonic()
        while True:
            for question_index, question in enumerate(questions):
                correct_ids = answers.get(question["id"])
                if correct_ids is None:
                    correct_ids = self.bank.correct_option_ids(question)
                if question["id"] in wrong_ids:
                    wrong = self._choose_wrong(question, correct_ids)
                    if wrong is None:
                        raise ValueError("无法为目标分数构造确定的错误答案")
                    selected = {wrong}
                else:
                    selected = correct_ids
                is_last = (
                    activity_type in (3, 4)
                    and question_index == len(questions) - 1
                    and duration_seconds is not None
                )
                if is_last:
                    remaining = duration_seconds - (time.monotonic() - started_at)
                    if remaining > 0:
                        time.sleep(remaining)
                now = time.monotonic()
                self._record_answer(
                    activity_type,
                    activity_id,
                    question,
                    selected,
                    current_index,
                    int(now - last_record + 0.999),
                )
                last_record = now
                processed += 1
                if processed % 25 == 0:
                    self.update(activity["name"], "进行中", f"{processed} 题", "答案已记录", "", 3)
            if activity_type in (3, 4) or int(paper.get("hasNext") or 2) != 1:
                break
            next_paper = self._post(
                PRACTICE_API + "/getNear",
                {"userActivityId": activity_id, "currentIndex": current_index, "nearType": 2},
            )
            next_index = int(next_paper.get("currentIndex") or current_index)
            if next_index == current_index:
                raise ValueError("练习分页没有前进；已停止，未交卷")
            current_index = next_index
            paper = next_paper
            questions = paper.get("questionList") or []
            if not questions:
                raise ValueError("练习下一页没有题目；已停止，未交卷")
        result = self._post(base + "/submitPaper", {"userActivityId": activity_id})
        history = self._post(base + "/listHistory", {"userActivityId": activity_id}) or []
        latest = history[0] if history else {}
        actual_score = latest.get("score", result.get("score", "未返回"))
        actual_time = latest.get("useTime", result.get("useTime", "未返回"))
        self.update(
            activity["name"], "已交卷", f"本次记录 {processed} 题",
            f"分数 {actual_score}；用时 {actual_time} 秒",
            "服务器成绩为准", 4,
        )
        if target_score is not None and actual_score != target_score:
            raise ValueError(f"实际得分 {actual_score} 与目标 {target_score} 不同")
        return result

    def work(self, selection: set[int], execute: bool, target_score: int | None, duration_seconds: int | None, seed: int) -> None:
        for activity_type in (1, 2):
            if activity_type in selection:
                self.run_activity(
                    activity_type, execute=execute,
                    target_score=target_score if len(selection) == 1 else None,
                    duration_seconds=duration_seconds if len(selection) == 1 else None,
                    seed=seed,
                )

    def exm(self, selection: set[int], execute: bool, target_score: int | None, duration_seconds: int | None, seed: int) -> None:
        for activity_type in (3, 4):
            if activity_type in selection:
                self.run_activity(
                    activity_type, execute=execute,
                    target_score=target_score if len(selection) == 1 else None,
                    duration_seconds=duration_seconds if len(selection) == 1 else None,
                    seed=seed,
                )


def apiMain(
    t: str, s: str, u: str, p: str, proxy: str | None, t_name: str,
    c: str, c_id: str, k_session: requests.Session,
    *, activity: str = "list", execute: bool = False,
    target_score: int | None = None, duration_seconds: int | None = None,
    seed: int = 1, bank_path: Path | None = None,
    client: API | None = None,
    on_authenticated: Callable[[API], None] | None = None,
) -> API:
    selected_bank = (bank_path or bundled_bank_path()).resolve()
    if client is None:
        api = API(t, s, u, p, proxy, t_name, k_session, selected_bank)
        api.setHeader()
        if not api.login():
            api.update("", "异常", "无", "登录失败", "", 1)
            raise RuntimeError("登录失败；请检查账号、密码、学校或验证码识别日志")
        api.update("", "进行中", "无", "登录成功", "", 1)
        if on_authenticated is not None:
            on_authenticated(api)
    else:
        if client.u != u or client.s != s or client.session is not k_session:
            raise ValueError("账号或学校已变更，请重新登录")
        api = client
        if api.bank_path != selected_bank:
            api.bank = AnswerResolver(selected_bank)
            api.bank_path = selected_bank
        api.update("", "进行中", "无", "复用已登录会话", "", 1)
    api.setInfo(c, c_id)
    if not api.get_coursesList():
        api.update("", "异常", "无", "获取项目失败", "", 1)
        raise RuntimeError("获取项目失败；请检查平台响应日志")
    if not api.getCourses():
        api.update("", "异常", "无", "找不到匹配项目", "", 1)
        raise RuntimeError("找不到匹配项目")
    api.list_activities()
    if activity == "list":
        return api
    # 桌面版的 auto 只执行尚无答题记录的活动，避免重复消耗考试次数。
    selection = {1, 2, 3} if activity == "all" else (
        {1, 2, 3, 4} if activity == "auto" else {ACTIVITY_TYPES[activity]}
    )
    for course in api.courseInfo:
        api.getDetail(course)
        api.handle()
        if activity == "auto":
            defaults = {1: (None, None), 2: (None, None), 3: (80, 240), 4: (90, 300)}
            for activity_type in (1, 2, 3, 4):
                item = api.activities.get(activity_type)
                if item is None:
                    continue
                if int((item.get("examData") or {}).get("doneTimes") or 0) > 0:
                    print(f"{item['name']}已有答题记录，自动模式跳过")
                    continue
                score, seconds = defaults[activity_type]
                api.run_activity(activity_type, execute=execute, target_score=score,
                                 duration_seconds=seconds, seed=seed)
        else:
            api.work(selection, execute, target_score, duration_seconds, seed)
            api.exm(selection, execute, target_score, duration_seconds, seed)
        if execute and not api.get_coursesList():
            raise ValueError("交卷后无法刷新项目进度")
        api.get_jd()
    api.update("", "已完成", api.jds["jd"], "任务完成", api.jds["jds"], 1)
    return api


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--username", default=os.getenv("MARS_USERNAME", "2602020150"))
    parser.add_argument("--school", default="江苏财经职业技术学院")
    parser.add_argument("--proxy", default=os.getenv("MARS_PROXY", urllib.request.getproxies().get("https")))
    parser.add_argument("--bank", type=Path, default=bundled_bank_path())
    parser.add_argument("--course", default="")
    parser.add_argument("--course-id", default="")
    parser.add_argument("--activity", choices=["list", *ACTIVITY_TYPES, "all"], default="list",
                        help="all 仅执行顺序、随机、模拟；正式考试须单独选 formal")
    parser.add_argument("--execute", action="store_true", help="允许创建答卷、记录答案并交卷")
    parser.add_argument("--target-score", type=int, help="模拟/正式考试目标分数；仅单个活动时使用")
    parser.add_argument("--duration-seconds", type=int,
                        help="模拟/正式考试目标用时秒数；网络延迟可能带来数秒偏差")
    parser.add_argument("--seed", type=int, default=1, help="选择故意答错题目的随机种子")
    args = parser.parse_args()
    if args.execute and args.activity == "list":
        parser.error("--execute 需要指定 --activity")
    if (args.target_score is not None or args.duration_seconds is not None) and args.activity not in {"mock", "formal"}:
        parser.error("控分控时仅适用于单个模拟考试或正式考试")
    password = os.getenv("MARS_PASSWORD") or getpass.getpass("竞赛平台密码: ")
    session = requests.Session()
    try:
        apiMain(
            "江苏省第十五届大学生安全知识竞赛", args.school,
            args.username, password, args.proxy, "deimos",
            args.course, args.course_id, session,
            activity=args.activity, execute=args.execute,
            target_score=args.target_score,
            duration_seconds=args.duration_seconds, seed=args.seed,
            bank_path=args.bank,
        )
    finally:
        session.close()


if __name__ == "__main__":
    main()
